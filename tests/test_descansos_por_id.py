"""Pruebas del cargador y el motor: descansos exclusivos de un ID."""

import gc
import tempfile
import unittest
from datetime import date, time
from pathlib import Path

import pandas as pd
import openpyxl

from core import COLOR_MAP, calcular_metricas, procesar_asistencias
from export import exportar_maestro_consolidado, exportar_reporte_empleado
from ingestion import cargar_plantilla, extraer_reglas
from utils import exportar_estadisticas_excel


class DescansosPorIdTests(unittest.TestCase):
    def procesar(self, descansos, marcas_extra=(), grupo="", rotaciones=(), festivos=()):
        catalogo = pd.DataFrame({
            "ID": ["123", "456"],
            "Nombre completo": ["Persona 123", "Persona 456"],
            "Estatus": ["Activo", "Activo"],
            "Tipo de personal": ["Interno", "Interno"],
            "Periodo ingreso": ["A", "A"],
        })
        rol = pd.DataFrame({"ID": ["123", "456"], "Rotación": [grupo, ""]})
        incidencias = pd.DataFrame(columns=["ID", "Ausencia Justificada", "Fecha Inicio", "Fecha Termino"])
        with tempfile.TemporaryDirectory() as temporal:
            plantilla = Path(temporal) / "plantilla.xlsx"
            with pd.ExcelWriter(plantilla, engine="openpyxl") as writer:
                catalogo.to_excel(writer, sheet_name="1_Catalogo_Personal", index=False)
                rol.to_excel(writer, sheet_name="2_Rol_Guardias", index=False)
                for hoja in ("3_Registro_Incidencias", "4_Vacaciones"):
                    incidencias.to_excel(writer, sheet_name=hoja, index=False)
                pd.DataFrame(rotaciones, columns=incidencias.columns).to_excel(
                    writer, sheet_name="5_Rotaciones", index=False)
                pd.DataFrame(festivos, columns=["Motivo", "Fecha"]).to_excel(
                    writer, sheet_name="6_Dias_Festivos", index=False)
                pd.DataFrame(descansos, columns=["ID", "Dia", "Fecha(Opcional)"]).to_excel(writer, sheet_name="7_Descansos", index=False)
            datos = cargar_plantilla(str(plantilla))
            marcas = [
                {"ID_Biometrico": id_emp, "Fecha": date(2026, 4, 1), "Hora_CheckIn": time(7), "Hora_CheckOut": time(17)}
                for id_emp in ("123", "456")
            ]
            datos["scanner"] = pd.DataFrame(marcas + list(marcas_extra))
            resultado = procesar_asistencias(datos, extraer_reglas(pd.DataFrame()))
            # Libera los lectores de Excel antes de retirar el archivo temporal en Windows.
            gc.collect()
            return resultado

    def dia(self, resultado, id_emp, fecha):
        return resultado[(resultado["ID"] == id_emp) & (resultado["Fecha"] == pd.Timestamp(fecha))].iloc[0]

    def test_descanso_semanal_solo_para_id_asignado_y_normalizado(self):
        resultado = self.procesar([("00123", "Domingo", None)])
        for fecha in ("2026-04-05", "2026-04-12", "2026-04-19", "2026-04-26"):
            asignado = self.dia(resultado, "123", fecha)
            self.assertEqual(asignado["Estatus"], "DESCANSO")
            self.assertEqual(asignado["Label"], "Descanso")
            self.assertEqual(asignado["Color_Hex"], COLOR_MAP["DESCANSO"]["hex"])
            self.assertEqual(self.dia(resultado, "456", fecha)["Estatus"], "NO_LABORABLE")
        self.assertEqual(calcular_metricas(resultado[resultado["ID"] == "123"])["descansos"], 4)
        self.assertEqual(calcular_metricas(resultado[resultado["ID"] == "456"])["descansos"], 0)

    def test_fecha_puntual_solo_en_fecha_e_id_indicados(self):
        resultado = self.procesar([("123.0", "Otro", date(2026, 4, 8))])
        self.assertEqual(self.dia(resultado, "123", "2026-04-08")["Estatus"], "DESCANSO")
        self.assertEqual(self.dia(resultado, "456", "2026-04-08")["Estatus"], "FALTA")
        self.assertEqual(self.dia(resultado, "123", "2026-04-07")["Estatus"], "FALTA")
        self.assertEqual(self.dia(resultado, "123", "2026-04-09")["Estatus"], "FALTA")
        self.assertEqual((resultado["Estatus"] == "DESCANSO").sum(), 1)

    def test_marca_en_descanso_individual_se_evalua_como_asistencia(self):
        resultado = self.procesar([("123", "Domingo", None)], marcas_extra=[{
            "ID_Biometrico": "123", "Fecha": date(2026, 4, 5),
            "Hora_CheckIn": time(7), "Hora_CheckOut": time(17),
        }])
        asignado = self.dia(resultado, "123", "2026-04-05")
        self.assertEqual(asignado["Estatus"], "ASISTENCIA")
        self.assertIn("Asistencia en Descanso", asignado["Notas"])
        self.assertEqual(self.dia(resultado, "123", "2026-04-12")["Estatus"], "DESCANSO")
        self.assertEqual(self.dia(resultado, "456", "2026-04-05")["Estatus"], "NO_LABORABLE")

    def test_generico_se_acumula_sin_extender_el_descanso_individual(self):
        resultado = self.procesar([("generico", "Sábado", None), ("123", "Domingo", None)])
        self.assertEqual(self.dia(resultado, "123", "2026-04-04")["Estatus"], "DESCANSO")
        self.assertEqual(self.dia(resultado, "456", "2026-04-04")["Estatus"], "DESCANSO")
        self.assertEqual(self.dia(resultado, "456", "2026-04-05")["Estatus"], "NO_LABORABLE")
        self.assertEqual(calcular_metricas(resultado[resultado["ID"] == "123"])["descansos"], 8)
        self.assertEqual(calcular_metricas(resultado[resultado["ID"] == "456"])["descansos"], 4)

    def test_guardia_conserva_prioridad_sobre_descanso_individual(self):
        resultado = self.procesar([("123", "Otro", date(2026, 4, 5))], grupo="A")
        asignado = self.dia(resultado, "123", "2026-04-05")
        self.assertEqual(asignado["Guardia_Tipo"], "Guardia")
        self.assertEqual(asignado["Estatus"], "FALTA_GUARDIA")
        self.assertEqual(self.dia(resultado, "123", "2026-04-04")["Estatus"], "NO_LABORABLE")
        self.assertEqual(self.dia(resultado, "456", "2026-04-05")["Estatus"], "NO_LABORABLE")

    def test_postguardia_asignada_en_domingo_conserva_su_estado(self):
        resultado = self.procesar([], grupo="D", marcas_extra=[
            {"ID_Biometrico": "123", "Fecha": date(2026, 4, 4),
             "Hora_CheckIn": time(7), "Hora_CheckOut": None},
            {"ID_Biometrico": "123", "Fecha": date(2026, 4, 5),
             "Hora_CheckIn": None, "Hora_CheckOut": time(7)},
        ])
        self.assertEqual(self.dia(resultado, "123", "2026-04-04")["Estatus"], "GUARDIA")
        self.assertEqual(self.dia(resultado, "123", "2026-04-05")["Estatus"], "POST_GUARDIA")
        self.assertEqual(self.dia(resultado, "456", "2026-04-05")["Estatus"], "NO_LABORABLE")

    def test_fin_de_semana_sin_asignacion_no_es_falta_aunque_haya_marca(self):
        resultado = self.procesar([], marcas_extra=[{
            "ID_Biometrico": "456", "Fecha": date(2026, 4, 4),
            "Hora_CheckIn": time(7), "Hora_CheckOut": time(17),
        }])
        for fecha in ("2026-04-04", "2026-04-05"):
            dia = self.dia(resultado, "456", fecha)
            self.assertEqual(dia["Estatus"], "NO_LABORABLE")
            self.assertEqual(dia["Color_Hex"], COLOR_MAP["NO_LABORABLE"]["hex"])
        self.assertIn("Marcas registradas", self.dia(resultado, "456", "2026-04-04")["Notas"])
        self.assertEqual(self.dia(resultado, "456", "2026-04-03")["Estatus"], "FALTA")
        metricas = calcular_metricas(resultado[resultado["ID"] == "456"])
        self.assertEqual(metricas["no_laborables"], 8)
        self.assertEqual(metricas["dias_evaluados"], 22)

    def test_rotacion_y_festivo_conservan_prioridad_en_fin_de_semana(self):
        resultado = self.procesar(
            [], rotaciones=[("456", "rotación", date(2026, 4, 4), date(2026, 4, 4))],
            festivos=[("Festivo de prueba", date(2026, 4, 5))],
        )
        self.assertEqual(self.dia(resultado, "456", "2026-04-04")["Estatus"], "ROTACION")
        self.assertEqual(self.dia(resultado, "456", "2026-04-05")["Estatus"], "FESTIVO")
        self.assertEqual(self.dia(resultado, "456", "2026-04-11")["Estatus"], "NO_LABORABLE")

    def test_excel_separa_no_laborable_y_festivo_y_pinta_la_leyenda(self):
        resultado = self.procesar([], festivos=[("Festivo de prueba", date(2026, 4, 5))])
        persona = resultado[resultado["ID"] == "456"]
        with tempfile.TemporaryDirectory() as temporal:
            rutas = [
                exportar_reporte_empleado(persona, temporal),
                exportar_maestro_consolidado(resultado, temporal),
                exportar_estadisticas_excel(resultado, temporal),
            ]
            libros = [openpyxl.load_workbook(ruta) for ruta in rutas]
            try:
                individual, maestro, estadisticas = libros
                hoja_ind = individual.active
                self.assertEqual(hoja_ind.cell(16, 5).value, "NL")  # sábado 4
                self.assertEqual(hoja_ind.cell(16, 6).value, "Fe")  # domingo 5
                self.assertEqual(hoja_ind.cell(16, 5).fill.fgColor.rgb, "FFD9D9D9")
                self.assertEqual(hoja_ind.cell(16, 6).fill.fgColor.rgb, "FFADD8E6")
                colores_ind = {
                    cell.value: cell.fill.fgColor.rgb
                    for row in hoja_ind.iter_rows(min_row=19, max_row=20) for cell in row
                    if cell.value in ("NL", "Fe")
                }
                self.assertEqual(colores_ind, {"NL": "FFD9D9D9", "Fe": "FFADD8E6"})
                self.assertTrue(all(
                    cell.fill.patternType == "solid"
                    for row in hoja_ind.iter_rows(min_row=19, max_row=20) for cell in row
                    if cell.value in ("As", "Re", "Fa", "Gu", "PG", "FG", "NL", "De", "Va", "Fe", "In", "Pe", "Co", "Ro")
                ))

                hoja_maestra = maestro.worksheets[0]
                self.assertEqual(hoja_maestra.column_dimensions["B"].width, 28)
                self.assertEqual(hoja_maestra.cell(5, 10).value, "NL")
                self.assertEqual(hoja_maestra.cell(5, 11).value, "Fe")
                self.assertEqual(hoja_maestra.cell(5, 10).fill.fgColor.rgb, "FFD9D9D9")
                self.assertEqual(hoja_maestra.cell(5, 11).fill.fgColor.rgb, "FFADD8E6")
                colores_maestro = {
                    hoja_maestra.cell(row, 1).value: hoja_maestra.cell(row, 1).fill.fgColor.rgb
                    for row in range(8, 22)
                }
                self.assertEqual(colores_maestro["NL"], "FFD9D9D9")
                self.assertEqual(colores_maestro["Fe"], "FFADD8E6")
                resumen = maestro["Resumen General"]
                resumen_por_estado = {
                    resumen.cell(fila, 1).value: resumen.cell(fila, 2).value
                    for fila in range(5, resumen.max_row + 1)
                }
                self.assertEqual(resumen_por_estado["No Laborable"], 14)
                self.assertEqual(resumen_por_estado["Festivo"], 2)

                hoja_stats = estadisticas["Resumen_Individual"]
                encabezados = [cell.value for cell in hoja_stats[2]]
                filas = [dict(zip(encabezados, [cell.value for cell in row]))
                         for row in hoja_stats.iter_rows(min_row=3, max_row=4)]
                persona_stats = next(row for row in filas if row["Nombre"] == "Persona 456")
                self.assertEqual(persona_stats["No laborables"], 7)
                self.assertEqual(persona_stats["Festivos"], 1)
            finally:
                for libro in libros:
                    libro.close()


if __name__ == "__main__":
    unittest.main()
